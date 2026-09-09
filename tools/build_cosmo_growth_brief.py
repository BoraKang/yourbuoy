from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

OUT = Path("deliverables/COSMO_매출_성장_구조_브리프.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "17365D"
BLUE = "2E74B5"
PALE = "EAF1F8"
LIGHT = "F3F5F7"
GRAY = "5B6573"
WHITE = "FFFFFF"
INK = "1F2933"
FONT = "Arial Unicode MS"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def font_run(run, size=10.5, bold=False, color=INK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_para(container, text="", size=10.5, bold=False, color=INK, before=0, after=6,
             align=None, keep=False):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    font_run(r, size=size, bold=bold, color=color)
    return p


def add_bullet(container, text, level=0, after=3):
    p = container.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    for r in p.runs:
        font_run(r, size=10.3)
    if not p.runs:
        font_run(p.add_run(text), size=10.3)
    else:
        p.runs[0].text = text
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    font_run(r, size=15 if level == 1 else 12.5, bold=True, color=BLUE if level == 1 else NAVY)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
for name, size, color, before, after in [
    ("Heading 1", 15, BLUE, 15, 7),
    ("Heading 2", 12.5, NAVY, 10, 5),
]:
    s = styles[name]
    s.font.name = FONT
    s._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font_run(hp.add_run("COSMO BUSINESS GROWTH BRIEF"), size=8.5, bold=True, color=GRAY)
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
font_run(fp.add_run("모드하우스 COSMO | 고객사 공유용"), size=8, color=GRAY)

add_para(doc, "BUSINESS GROWTH BRIEF", size=9, bold=True, color=BLUE, after=7)
add_para(doc, "COSMO 매출 성장 구조 및 조직 운영 관점", size=24, bold=True, color=NAVY, after=5, keep=True)
add_para(doc, "세 가지 매출 성장 축과 Tribe 구성 가능성", size=12, color=GRAY, after=15)

callout = doc.add_table(rows=1, cols=1)
callout.alignment = WD_TABLE_ALIGNMENT.LEFT
set_table_geometry(callout, [9360])
cell = callout.cell(0, 0)
set_cell_shading(cell, PALE)
p = cell.paragraphs[0]
p.paragraph_format.space_after = Pt(2)
font_run(p.add_run("핵심 요약"), size=10, bold=True, color=BLUE)
add_bullet(cell, "COSMO의 최상위 목표는 월간 플랫폼 총매출 증대")
add_bullet(cell, "Level 2는 Objekt·멤버십·아티스트 상품의 3개 병렬 성장 축")
add_bullet(cell, "세 축은 조직 조건을 충족할 경우 각각의 Tribe로 운영 가능")
add_bullet(cell, "고객 성장은 일방향 퍼널이 아니라 참여와 구매가 반복되는 플라이휠 구조")

heading(doc, "1. 목표 및 지표 구조")
add_para(doc, "Level 1 | 최상위 목표", size=11, bold=True, color=NAVY, after=3)
add_para(doc, "월간 플랫폼 총매출 증대", size=13, bold=True, color=INK, after=4)
add_para(doc, "COSMO에서 발생하는 플랫폼 기반 매출의 합계", color=GRAY, after=8)
add_para(doc, "총매출 = Objekt 매출 + 멤버십 매출 + 아티스트 상품 매출", size=10.7, bold=True, color=BLUE, after=10)

add_para(doc, "Level 2 | 세 가지 매출 성장 축", size=11, bold=True, color=NAVY, after=6)
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.style = "Table Grid"
headers = ["성장 축", "정의", "핵심 지표", "주요 실행 과제"]
for i, txt in enumerate(headers):
    c = table.rows[0].cells[i]
    set_cell_shading(c, NAVY)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run(txt), size=9.3, bold=True, color=WHITE)

rows = [
    ("Objekt 매출", "디지털 포토카드 및 번들 판매 매출", "활성 수집가 수\n1인당 구매 횟수\n평균 판매가격", "한정 Objekt 출시\n첫 구매 전환\n컬렉션 완성 보상"),
    ("멤버십 매출", "아티스트별·통합 유료 멤버십의 반복 매출", "유료 멤버십 수\n유료 전환율\n이탈률", "멀티 아티스트 상품\n갱신·혜택 강화\n발매 연계 프로모션"),
    ("아티스트 상품 매출", "앨범·공연 티켓·MD 등 관련 상품 매출", "구매 건수\n평균 주문금액\n이벤트 추가 매출", "이벤트 한정 판매\n추천 보상\n개인화 추천"),
]
for ridx, row in enumerate(rows):
    cells = table.add_row().cells
    for i, txt in enumerate(row):
        if ridx % 2 == 1:
            set_cell_shading(cells[i], LIGHT)
        p = cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        for j, line in enumerate(txt.split("\n")):
            if j:
                p.add_run("\n")
            font_run(p.add_run(line), size=9.1, bold=(i == 0), color=NAVY if i == 0 else INK)
set_table_geometry(table, [1680, 2680, 2350, 2650])

heading(doc, "2. 세 성장 축을 3개 Tribe로 볼 수 있는가")
add_para(doc, "결론 | 조직 요건을 충족한다면 3개 Tribe로 운영 가능", size=11.3, bold=True, color=BLUE, after=5)
add_para(doc, "현재 지표 구조만으로는 ‘3개 Tribe’라기보다 ‘3개 성장 축’ 또는 ‘3개 수익 스트림’으로 정의하는 것이 정확함", after=8)

tribe_table = doc.add_table(rows=1, cols=3)
tribe_table.style = "Table Grid"
tribe_table.alignment = WD_TABLE_ALIGNMENT.LEFT
for i, txt in enumerate(["조직 단위", "핵심 미션", "대표 KPI"]):
    c = tribe_table.rows[0].cells[i]
    set_cell_shading(c, BLUE)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run(txt), size=9.5, bold=True, color=WHITE)
for row in [
    ("Objekt Tribe", "수집가 확대 및 반복 구매 증대", "활성 수집가·구매 빈도·객단가"),
    ("Membership Tribe", "유료 전환 및 구독 유지", "유료 회원·전환율·이탈률"),
    ("Artist Commerce Tribe", "아티스트 IP 상품 구매 확대", "구매 건수·주문금액·이벤트 매출"),
]:
    cells = tribe_table.add_row().cells
    for i, txt in enumerate(row):
        p = cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        font_run(p.add_run(txt), size=9.4, bold=(i == 0), color=NAVY if i == 0 else INK)
set_table_geometry(tribe_table, [2260, 3980, 3120])

add_para(doc, "Tribe 성립 조건", size=11, bold=True, color=NAVY, before=8, after=4)
for item in [
    "고유한 사업 목표와 KPI를 보유",
    "기획·디자인·개발·마케팅 등 실행 역량을 내부에 확보",
    "문제 발견부터 실행·성과까지 독립적으로 책임",
    "프로젝트 종료와 무관하게 장기적으로 유지되는 조직",
]:
    add_bullet(doc, item)
add_para(doc, "위 조건을 충족하지 못하고 프로젝트별 인력 공유가 중심이라면 ‘Tribe’보다 ‘Growth Pillar’ 또는 ‘Workstream’이 적합", size=9.8, color=GRAY, before=3, after=6)

heading(doc, "3. 핵심 성장 구조")
add_para(doc, "매출 지표 구조 | 병렬형", size=11, bold=True, color=NAVY, after=3)
add_para(doc, "Objekt·멤버십·아티스트 상품 매출은 특정 순서로 발생하지 않으며, 세 축이 동시에 월간 플랫폼 총매출에 기여", after=7)

flow = doc.add_table(rows=1, cols=3)
flow.alignment = WD_TABLE_ALIGNMENT.LEFT
for i, txt in enumerate(["Objekt 매출", "멤버십 매출", "아티스트 상품 매출"]):
    c = flow.rows[0].cells[i]
    set_cell_shading(c, PALE if i != 1 else "DCEAF7")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    font_run(p.add_run(txt), size=10.2, bold=True, color=NAVY)
set_table_geometry(flow, [3120, 3120, 3120])
add_para(doc, "▼ 세 매출 축의 합계", size=9.3, bold=True, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=2)
add_para(doc, "월간 플랫폼 총매출", size=13, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

add_para(doc, "고객 성장 구조 | 순환형 플라이휠", size=11, bold=True, color=NAVY, after=5)
wheel = doc.add_table(rows=1, cols=5)
wheel.alignment = WD_TABLE_ALIGNMENT.LEFT
steps = ["팬 유입", "콘텐츠·투표 참여", "수집·구독·상품 구매", "팬 경험·관계 강화", "재방문·재구매·추천"]
for i, txt in enumerate(steps):
    c = wheel.rows[0].cells[i]
    set_cell_shading(c, NAVY if i in (0, 4) else BLUE)
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    font_run(p.add_run(txt), size=8.8, bold=True, color=WHITE)
set_table_geometry(wheel, [1450, 1950, 2220, 1870, 1870])
add_para(doc, "재방문·재구매·추천이 다시 신규 팬 유입과 참여 확대로 연결", size=9.5, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, before=4, after=9)

heading(doc, "4. 권장 표현")
add_para(doc, "“COSMO는 Objekt, 멤버십, 아티스트 상품이라는 세 가지 매출 성장 축을 중심으로 운영된다. 각 축은 독립적인 수익 지표를 보유하지만, 팬의 참여·수집·구독·구매가 상호 강화되는 플라이휠을 형성한다. 조직적으로는 각 축에 독립적인 목표와 실행 권한을 부여할 경우 3개 Tribe 체계로 운영할 수 있다.”", size=10.7, bold=True, color=NAVY, before=2, after=4)

doc.save(OUT)
print(OUT.resolve())
